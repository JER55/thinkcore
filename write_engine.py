#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Scrive i turni generati direttamente nel foglio Turni di TOOL_LIDL.xlsx,
modificando SOLO l'XML del foglio Turni dentro lo zip .xlsx, lasciando
ogni altro file (grafici, data bar, formattazione condizionale, altri
fogli) byte-per-byte identico all'originale.

Perché non openpyxl.load_workbook + save(): verificato che questo causa
perdita reale di regole di formattazione condizionale avanzate (data bar)
su Dashboard e OreRichieste (da 6 a 3 regole) — vedi test in questa sessione.

v2 (fix):
- Trova il file XML del foglio "Turni" leggendo workbook.xml + rels,
  invece di assumere "sheet5.xml" (che dipende dall'ordine dei fogli).
- Se una cella target non esiste (Excel omette celle vuote senza stile),
  la inserisce nella posizione corretta invece di fallire.
- Stile di fallback preso dalla colonna, non più un numero magico.
"""
import re
import zipfile
from collections import Counter
from openpyxl.utils import get_column_letter

DAY_PAIRS = [(6, 7), (9, 10), (12, 13), (15, 16), (18, 19), (21, 22), (24, 25)]  # F,G | I,J | ...


# ---------- helper: trova il file XML del foglio per nome ----------
def _find_sheet_file(zin, sheet_name):
    wb_xml = zin.read('xl/workbook.xml').decode('utf-8')
    # Prova entrambi gli ordini di attributi (Excel non li garantisce)
    m = re.search(r'<sheet\b[^>]*\bname="' + re.escape(sheet_name) + r'"[^>]*\br:id="([^"]+)"', wb_xml)
    if not m:
        m = re.search(r'<sheet\b[^>]*\br:id="([^"]+)"[^>]*\bname="' + re.escape(sheet_name) + r'"', wb_xml)
    if not m:
        raise RuntimeError(f'Foglio "{sheet_name}" non trovato in xl/workbook.xml. '
                           f'Verifica che il modello contenga un foglio con questo nome esatto.')
    rid = m.group(1)

    rels_xml = zin.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    m = re.search(r'<Relationship\b[^>]*\bId="' + re.escape(rid) + r'"[^>]*\bTarget="([^"]+)"', rels_xml)
    if not m:
        m = re.search(r'<Relationship\b[^>]*\bTarget="([^"]+)"[^>]*\bId="' + re.escape(rid) + r'"', rels_xml)
    if not m:
        raise RuntimeError(f'Relationship {rid} non trovata in xl/_rels/workbook.xml.rels.')
    target = m.group(1)
    if target.startswith('/'):
        return target.lstrip('/')
    return 'xl/' + target


# ---------- helper: colonna letterale -> numero ----------
def _col_num(letter):
    n = 0
    for ch in letter:
        n = n * 26 + (ord(ch) - ord('A') + 1)
    return n


# ---------- helper: stile di una cella esistente nella riga ----------
def _style_of(row_xml, ref):
    m = re.search(r'<c\s+r="' + re.escape(ref) + r'"[^>]*?\bs="(\d+)"', row_xml)
    return m.group(1) if m else None


# ---------- helper: stile di fallback per una colonna ----------
def _column_default_style(sheet_xml, col_letter):
    styles = re.findall(r'<c\s+r="' + re.escape(col_letter) + r'\d+"[^>]*?\bs="(\d+)"', sheet_xml)
    if not styles:
        return None
    return Counter(styles).most_common(1)[0][0]


# ---------- helper: inserisci o sostituisci una cella dentro la riga ----------
def _put_cell(row_xml, ref, new_cell):
    # 1) Sostituzione se la cella è presente (self-closing o con contenuto)
    pat = (r'<c\s+r="' + re.escape(ref) + r'"(?:\s[^>]*?)?/>'
           r'|<c\s+r="' + re.escape(ref) + r'"(?:\s[^>]*?)?>.*?</c>')
    new_row, n = re.subn(pat, new_cell, row_xml, count=1, flags=re.DOTALL)
    if n == 1:
        return new_row

    # 2) Inserimento in ordine di colonna
    target = _col_num(re.match(r'([A-Z]+)', ref).group(1))
    cells = list(re.finditer(r'<c\s+r="([A-Z]+\d+)"', row_xml))
    insert_at = None
    for mm in cells:
        if _col_num(re.match(r'([A-Z]+)', mm.group(1)).group(1)) > target:
            insert_at = mm.start()
            break
    if insert_at is None:
        m_close = re.search(r'</row>', row_xml)
        if not m_close:
            raise RuntimeError(f'Impossibile inserire {ref}: </row> non trovato.')
        insert_at = m_close.start()
    return row_xml[:insert_at] + new_cell + row_xml[insert_at:]


def write_shifts_surgical(tool_lidl_path: str, out_path: str,
                          EMP: list, shifts: dict, day_label: dict = None,
                          max_slots: int = 22, first_row: int = 5) -> dict:
    """
    EMP: lista dipendenti presenti (stesso ordine di Personale, come li
         costruisce generatore_turni_v6.py) — EMP[i]['id'] indicizza shifts.
    shifts: dict {emp_id: [ (start,end)|None per 7 giorni ]}
    day_label: opzionale, dict {emp_id: ['R'|'P'|None per 7 giorni]} (attributo
               DAY_LABEL del motore) — distingue Riposo da Permesso nei
               giorni senza turno. Se assente, tutti i giorni None -> 'R'.
    Scrive SOLO nelle colonne F,G,I,J,L,M,O,P,R,S,U,V,X,Y (In/Out per giorno),
    righe first_row..first_row+max_slots-1. Non tocca A:E (nomi/ruolo,
    self-updating) né le colonne Ore/Tot/Delta (formule).
    """
    with zipfile.ZipFile(tool_lidl_path, 'r') as zin:
        sheet_file = _find_sheet_file(zin, 'Turni')
        sheet_xml = zin.read(sheet_file).decode('utf-8')

    # Trova le righe target con le loro posizioni nel foglio. Processiamo
    # dalla più a valle alla più a monte: gli splice successivi non
    # invalidano gli indici delle righe ancora da modificare.
    row_matches = []
    for i in range(max_slots):
        rn = first_row + i
        m = re.search(r'<row\b[^>]*\br="' + str(rn) + r'"[^>]*?(?:/>|>.*?</row>)',
                      sheet_xml, flags=re.DOTALL)
        if not m:
            # Riga assente nell'XML (slot non usato): saltiamo.
            continue
        row_matches.append((m.start(), m.end(), rn, i, m.group(0)))

    n_written = 0
    for start, end, rn, i, row_xml in sorted(row_matches, key=lambda x: -x[0]):
        # Se la riga è self-closing (<row .../>), aprila per poterci scrivere
        if '</row>' not in row_xml:
            row_xml = row_xml.rstrip()[:-2].rstrip() + '></row>'

        if i < len(EMP):
            e = EMP[i]
            day_shifts = shifts[e['id']]
            labels = (day_label or {}).get(e['id'], [None] * 7)
        else:
            day_shifts = [None] * 7   # slot inutilizzato -> 'R' ovunque
            labels = [None] * 7

        for d, (ci, co) in enumerate(DAY_PAIRS):
            s = day_shifts[d]
            for col, is_start in ((ci, True), (co, False)):
                ref = f"{get_column_letter(col)}{rn}"
                style = _style_of(row_xml, ref)
                if style is None:
                    style = _column_default_style(sheet_xml, get_column_letter(col))
                s_attr = f' s="{style}"' if style is not None else ''

                if s is None:
                    # Riposo/permesso come inlineStr: zero dipendenza da
                    # sharedStrings.xml (che è file-specifico e si sposta a
                    # ogni modifica in Excel). Byte-preservation invariata.
                    label = 'P' if labels[d] == 'P' else 'R'
                    new_cell = f'<c r="{ref}"{s_attr} t="inlineStr"><is><t>{label}</t></is></c>'
                else:
                    val = (s[0] if is_start else s[1]) / 24.0
                    new_cell = f'<c r="{ref}"{s_attr} t="n"><v>{val!r}</v></c>'

                row_xml = _put_cell(row_xml, ref, new_cell)
                n_written += 1

        # Sostituisci la riga nel foglio
        sheet_xml = sheet_xml[:start] + row_xml + sheet_xml[end:]

    # Riscrive lo zip: tutti i file identici tranne il foglio Turni
    with zipfile.ZipFile(tool_lidl_path, 'r') as zin, \
         zipfile.ZipFile(out_path, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == sheet_file:
                data = sheet_xml.encode('utf-8')
            zout.writestr(item, data)

    return {'celle_scritte': n_written, 'foglio': sheet_file}